import io

import docx
from pypdf import PdfReader

from app.services import tailor as svc
from tests.conftest import register
from tests.test_readiness import create_target

RESUME = ("Jane Doe\njane@example.com | 555-0100 | Austin, TX\nEXPERIENCE\n"
          "Software Engineer, Acme Corp, 2021-2024\n"
          "- Worked on the payments service using Python and Django, handling 2,000 requests per second.\n"
          "EDUCATION\nB.S. Computer Science, University of Texas, 2021\n")
JD = "We are hiring a backend engineer with Python, Django, PostgreSQL and AWS experience. " * 2


def content(company="Acme Corp", bullet="Built the payments service in Python/Django serving 2000 req/s and "
                                         "cut latency by [X%]", cert=()):
    return {
        "contact": {"name": "Jane Doe", "email": "jane@example.com", "phone": "555-0100",
                    "location": "Austin, TX", "links": ["https://github.com/jane"]},
        "headline": "Backend Engineer | Python, Django",
        "summary": "Backend engineer with 3 years building payment systems.",
        "skills": [{"category": "Languages", "items": ["Python", "SQL"]}],
        "experience": [{"title": "Software Engineer", "company": company, "location": "Austin, TX",
                        "start": "2021", "end": "2024", "bullets": [bullet]}],
        "projects": [],
        "education": [{"degree": "B.S. Computer Science", "institution": "University of Texas",
                       "location": "", "start": "", "end": "2021", "details": ""}],
        "certifications": list(cert),
    }


def llm_result(**kw):
    return {"resume": content(**kw),
            "changes": [{"section": "Experience", "before": "Worked on", "after": "Built", "reason": "verb"}],
            "keywords_added": ["Python", "Django"], "keywords_missing": ["AWS"]}


def files():
    return {"resume": ("cv.txt", io.BytesIO(RESUME.encode()), "text/plain")}


def create(client, fake_llm, data=None, **kw):
    fake_llm.push(llm_result(**kw))
    data = data or {"job_description": JD, "job_title": "Backend"}
    return client.post("/api/tailor", files=files(), data=data)


def test_fact_check_flags_invented_facts():
    ok = svc.fact_check(RESUME, svc.ResumeContent.model_validate(content()))
    # "2000" matches "2,000"; "3" (years) is new and flagged; placeholders are ignored.
    assert ok == ['The figure "3" is not in your original resume. Make sure it is accurate before sending.']
    bad = svc.fact_check(RESUME, svc.ResumeContent.model_validate(
        content(company="Google", bullet="Grew revenue 45%", cert=["AWS Certified Developer"])))
    assert any("Google" in w for w in bad)
    assert any("45%" in w for w in bad)
    assert any("AWS Certified Developer" in w for w in bad)


def test_create_requires_job_description_without_charging(client, user, fake_llm):
    r = client.post("/api/tailor", files=files(), data={"job_title": "Backend"})
    assert r.status_code == 400
    assert client.get("/api/auth/me").json()["credits"] == 10
    assert fake_llm.calls == []


def test_create_edit_export_delete(client, user, fake_llm):
    r = create(client, fake_llm)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["placeholders"] == 1
    assert body["keywordsMissing"] == ["AWS"]
    assert client.get("/api/auth/me").json()["credits"] == 6
    rid = body["id"]

    edited = body["content"]
    edited["experience"][0]["bullets"][0] = "Built payments serving 2000 req/s, cutting latency 30%"
    r = client.put(f"/api/tailor/{rid}", json=edited)
    assert r.status_code == 200 and r.json()["placeholders"] == 0

    r = client.get(f"/api/tailor/{rid}/export", params={"format": "docx"})
    assert r.status_code == 200
    assert r.headers["content-disposition"].endswith('Jane_Doe_Backend.docx"')
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    for expected in ("Jane Doe", "EXPERIENCE", "cutting latency 30%", "University of Texas"):
        assert expected in text

    r = client.get(f"/api/tailor/{rid}/export", params={"format": "pdf"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    pdf_text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(r.content)).pages)
    assert "Jane Doe" in pdf_text and "EDUCATION" in pdf_text

    assert [x["id"] for x in client.get("/api/tailor").json()["resumes"]] == [rid]
    assert client.delete(f"/api/tailor/{rid}").status_code == 204
    assert client.get(f"/api/tailor/{rid}").status_code == 404


def test_pdf_handles_unicode():
    c = svc.ResumeContent.model_validate(content())
    c.contact.name = "José Müller–Ørsted 王"
    c.summary = "Built “resilient” systems — fast…"
    pdf = __import__("app.services.exporters", fromlist=["resume_pdf"]).resume_pdf(c)
    assert pdf.startswith(b"%PDF")


def test_uses_active_target_job_description(client, user, fake_llm):
    create_target(client, fake_llm, jd=JD)
    r = create(client, fake_llm, data={"use_target": "true"})
    assert r.status_code == 201, r.text
    assert r.json()["jobTitle"] == "Frontend Engineer"
    assert "PostgreSQL and AWS" in fake_llm.calls[-1][1]["prompt"]


def test_tailored_resumes_are_private(client, user, fake_llm):
    rid = create(client, fake_llm).json()["id"]
    client.post("/api/auth/logout")
    register(client, email="eve@example.com")
    assert client.get(f"/api/tailor/{rid}").status_code == 404
    assert client.get(f"/api/tailor/{rid}/export").status_code == 404
    assert client.delete(f"/api/tailor/{rid}").status_code == 404


def test_schema_is_strict_compatible():
    from openai.lib._parsing._responses import type_to_text_format_param

    assert type_to_text_format_param(svc.TailorLLM)["strict"] is True


def test_warnings_update_after_edits(client, user, fake_llm):
    body = create(client, fake_llm, company="Google", bullet="Grew revenue 45%").json()
    assert any("Google" in w for w in body["warnings"]) and any("45%" in w for w in body["warnings"])
    edited = body["content"]
    edited["experience"][0]["company"] = "Acme Corp"
    edited["experience"][0]["bullets"] = ["Grew revenue [X%]"]
    edited["summary"] = "Backend engineer building payment systems."
    after = client.put(f"/api/tailor/{body['id']}", json=edited).json()
    assert after["warnings"] == [] and after["placeholders"] == 1
