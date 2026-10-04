"""ATS-safe document exports (DOCX via python-docx, PDF via fpdf2).

Layout rules that keep applicant-tracking systems happy: one column, no tables or text boxes,
standard section headings, real text (not images), common fonts.
"""

from __future__ import annotations

import io
import unicodedata

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF

from .tailor import ResumeContent


def _dates(start: str, end: str) -> str:
    return " – ".join(p for p in (start.strip(), end.strip()) if p)


def _contact_line(content: ResumeContent) -> str:
    c = content.contact
    return " | ".join(p for p in (c.email, c.phone, c.location, *c.links) if p)


def _sections(content: ResumeContent):
    """Yield (heading, [(kind, text)]) in order; kind is 'line', 'sub', 'bullet' or 'meta'."""
    if content.summary.strip():
        yield "SUMMARY", [("line", content.summary.strip())]
    if content.skills:
        yield "SKILLS", [("line", f"{g.category}: {', '.join(g.items)}" if g.category else ", ".join(g.items))
                         for g in content.skills if g.items]
    if content.experience:
        items = []
        for r in content.experience:
            items.append(("sub", " — ".join(p for p in (r.title, r.company) if p)))
            meta = " | ".join(p for p in (r.location, _dates(r.start, r.end)) if p)
            if meta:
                items.append(("meta", meta))
            items += [("bullet", b) for b in r.bullets if b.strip()]
        yield "EXPERIENCE", items
    if content.projects:
        items = []
        for p in content.projects:
            title = p.name + (f" ({', '.join(p.tech)})" if p.tech else "")
            items.append(("sub", title))
            if p.link:
                items.append(("meta", p.link))
            items += [("bullet", b) for b in p.bullets if b.strip()]
        yield "PROJECTS", items
    if content.education:
        items = []
        for e in content.education:
            items.append(("sub", " — ".join(p for p in (e.degree, e.institution) if p)))
            meta = " | ".join(p for p in (e.location, _dates(e.start, e.end), e.details) if p)
            if meta:
                items.append(("meta", meta))
        yield "EDUCATION", items
    if content.certifications:
        yield "CERTIFICATIONS", [("bullet", c) for c in content.certifications if c.strip()]


# ---------------------------------------------------------------- DOCX
def resume_docx(content: ResumeContent) -> bytes:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(42)
        section.left_margin = section.right_margin = Pt(50)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(content.contact.name or "Your Name")
    run.bold = True
    run.font.size = Pt(18)
    for text in (content.headline, _contact_line(content)):
        if text:
            p = document.add_paragraph(text)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for heading, items in _sections(content):
        h = document.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        r = h.add_run(heading)
        r.bold = True
        r.font.size = Pt(12)
        for kind, text in items:
            if kind == "bullet":
                document.add_paragraph(text, style="List Bullet")
            else:
                p = document.add_paragraph()
                run = p.add_run(text)
                run.bold = kind == "sub"
                run.italic = kind == "meta"
                p.paragraph_format.space_after = Pt(2)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def letter_docx(body: str, subject: str = "") -> bytes:
    document = docx.Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    if subject:
        p = document.add_paragraph()
        p.add_run(f"Subject: {subject}").bold = True
    for para in body.split("\n\n"):
        document.add_paragraph(para.strip())
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
_REPLACEMENTS = {"–": "-", "—": "-", "‘": "'", "’": "'", "“": '"',
                 "”": '"', "•": "-", "…": "...", " ": " "}


def _latin1(text: str) -> str:
    """Core PDF fonts are Latin-1; map typographic characters and strip accents beyond that."""
    for k, v in _REPLACEMENTS.items():
        text = text.replace(k, v)
    try:
        text.encode("latin-1")
        return text
    except UnicodeEncodeError:
        normalized = unicodedata.normalize("NFKD", text)
        return normalized.encode("latin-1", "ignore").decode("latin-1")


def resume_pdf(content: ResumeContent) -> bytes:
    pdf = FPDF(format="Letter")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    width = pdf.epw

    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(width, 9, _latin1(content.contact.name or "Your Name"), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    for text in (content.headline, _contact_line(content)):
        if text:
            pdf.multi_cell(width, 5, _latin1(text), align="C", new_x="LMARGIN", new_y="NEXT")

    for heading, items in _sections(content):
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11.5)
        pdf.cell(width, 6, heading, new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
        pdf.ln(1.5)
        for kind, text in items:
            text = _latin1(text)
            if kind == "bullet":
                pdf.set_font("Helvetica", "", 10)
                pdf.set_x(pdf.l_margin + 3)
                pdf.multi_cell(width - 3, 5, f"- {text}", new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.set_font("Helvetica", "B" if kind == "sub" else ("I" if kind == "meta" else ""), 10)
                pdf.multi_cell(width, 5, text, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
